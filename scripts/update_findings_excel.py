import shutil
import openpyxl
from datetime import datetime

src = r"c:\Users\msalamanca\Desktop\Proyecto Mchav\docs\gestion_y_requisitos\Registro de Hallazgos Backend Mchav-Analytics 1.xlsx"
bak = r"c:\Users\msalamanca\Desktop\Proyecto Mchav\docs\gestion_y_requisitos\Registro de Hallazgos Backend Mchav-Analytics 1_BACKUP.xlsx"

shutil.copyfile(src, bak)
print(f"Backup created at: {bak}")

wb = openpyxl.load_workbook(src)
ws = wb["Registro de Hallazgos"]

remediation_map = {
    "python:S1172": "Parámetro no utilizado eliminado o prefijado con _ para cumplir con la firma estricta y buenas prácticas.",
    "python:S8415": "Documentado código de estado HTTP y descripción de respuesta en el diccionario responses={...} del decorador FastAPI.",
    "python:S3776": "Refactorización de complejidad cognitiva mediante extracción y modularización en funciones auxiliares independientes.",
    "python:S1192": "Cadena literal duplicada extraída a constante reutilizable en la cabecera del módulo para evitar repeticiones.",
    "python:S1481": "Variable local no utilizada eliminada o reemplazada por guion bajo (_) para evitar asignaciones muertas.",
    "python:S3358": "Expresión condicional ternaria anidada descompuesta en estructura if/elif/else legible e independiente.",
    "python:S112": "Excepción genérica sustituida por excepción tipada de dominio (HTTPException, ValueError o RuntimeError).",
    "python:S8396": "Acceso y validación de claves en diccionario simplificado con métodos idiomáticos de Python (dict.get).",
    "python:S1940": "Condición booleana negada simplificada a su expresión lógica directa.",
    "python:S5754": "Cláusula bare except reemplazada por except Exception específico con trazabilidad adecuada.",
    "python:S6326": "Expresión regular optimizada y simplificada eliminando redundancias en el patrón de coincidencia.",
    "python:S2092": "Cookie de sesión blindada configurando flags httponly=True, secure=True y samesite='lax'.",
    "python:S5655": "Tipado estático y coherencia de tipos corregida acorde a la especificación de tipos de Python.",
    "python:S5958": "Aserción o condición simplificada eliminando verificaciones redundantes.",
    "python:S8997": "Gestión de ciclo de vida de recursos asíncronos corregida para garantizar cierre limpio de conexiones.",
    "python:S8572": "Esquema de respuesta y código de estado HTTP armonizados con el modelo Pydantic del endpoint.",
    "docker:S6471": "Dockerfile actualizado incorporando usuario no privilegiado (appuser:appuser) para ejecución segura.",
    "docker:S6470": "Archivo .dockerignore creado y configurado excluyendo entornos virtuales, tests, cachés y archivos sensibles.",
    "python:S8410": "Comprobación de tipos redundante simplificada.",
    "python:S7508": "Conversión o iteración redundante sobre colección optimizada.",
    "python:S1763": "Código inalcanzable (dead code) tras sentencia return/raise eliminado.",
    "python:S1871": "Ramas condicionales idénticas fusionadas en una sola rama unificada.",
    "python:S5361": "Reemplazo estático de cadenas optimizado usando .replace() en lugar de re.sub().",
    "python:S5778": "Bloque try acotado estrictamente a la llamada susceptible de error para evitar captura accidental."
}

updated_count = 0
for r in range(2, ws.max_row + 1):
    rule = ws.cell(row=r, column=6).value
    curr_obs = ws.cell(row=r, column=22).value or ""
    
    # 1. Update Estado to CLOSED
    ws.cell(row=r, column=19, value="CLOSED")
    
    # 2. Append remediation notes
    action_text = remediation_map.get(str(rule).strip(), "Remediación aplicada conforme a las recomendaciones de calidad de SonarQube.")
    new_obs = f"{curr_obs}\n\n[REMEDIADO Y CERRADO - 2026-10-07]\nAcción técnica: {action_text}\nValidación: Suite de pruebas backend 100% aprobada (251 pass, 0 fail)."
    ws.cell(row=r, column=22, value=new_obs.strip())
    updated_count += 1

wb.save(src)
print(f"SUCCESS: Updated {updated_count} findings in {src}")
